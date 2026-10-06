import logging
import os
import re
import unicodedata
from datetime import datetime
from pathlib import Path
from tkinter import Tk, filedialog, messagebox

import pandas as pd
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURAZIONE
# ============================================================

APP_NAME = "Smart Sales Report"
VERSION = "4.5 PRO"
BASE_DIR = Path(__file__).resolve().parent
LOG_FILE = BASE_DIR / "smart_sales_report.log"

COLORE_PRINCIPALE = "17365D"
COLORE_SECONDARIO = "4472C4"
COLORE_CHIARO = "D9EAF7"
COLORE_VERDE = "E2F0D9"
COLORE_GIALLO = "FFF2CC"
COLORE_ROSSO = "FCE4D6"
COLORE_BIANCO = "FFFFFF"
COLORE_GRIGIO = "F2F2F2"

logging.basicConfig(
    filename=LOG_FILE,
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
    encoding="utf-8",
)
logger = logging.getLogger(__name__)


# ============================================================
# INTERFACCIA
# ============================================================

def crea_root():
    root = Tk()
    root.withdraw()
    return root


def scegli_file():
    root = crea_root()
    percorso = filedialog.askopenfilename(
        title="Seleziona il file del cliente",
        filetypes=[
            ("File supportati", "*.csv *.xlsx *.xlsm"),
            ("File CSV", "*.csv"),
            ("File Excel", "*.xlsx *.xlsm"),
        ],
    )
    root.destroy()
    return Path(percorso) if percorso else None


# ============================================================
# LETTURA FILE
# ============================================================

def leggi_csv(percorso):
    codifiche = ["utf-8-sig", "utf-8", "cp1252", "latin-1"]
    ultimo_errore = None

    for codifica in codifiche:
        try:
            df = pd.read_csv(
                percorso,
                encoding=codifica,
                sep=None,
                engine="python",
            )
            logger.info("CSV letto con codifica %s", codifica)
            return df
        except Exception as errore:
            ultimo_errore = errore

    raise ValueError(f"Impossibile leggere il file CSV: {ultimo_errore}")


def carica_dati(percorso):
    estensione = percorso.suffix.lower()

    if estensione == ".csv":
        return leggi_csv(percorso)

    if estensione in {".xlsx", ".xlsm"}:
        try:
            excel = pd.ExcelFile(percorso)
            if not excel.sheet_names:
                raise ValueError("Il file Excel non contiene fogli leggibili.")

            foglio = excel.sheet_names[0]
            logger.info("Excel: uso il primo foglio '%s'", foglio)
            return pd.read_excel(excel, sheet_name=foglio)
        except Exception as errore:
            raise ValueError(f"Impossibile leggere il file Excel: {errore}") from errore

    raise ValueError(f"Formato {estensione} non supportato.")


# ============================================================
# NORMALIZZAZIONE E RICONOSCIMENTO COLONNE
# ============================================================

def normalizza_testo(testo):
    testo = str(testo).strip().lower()
    testo = unicodedata.normalize("NFKD", testo)
    testo = "".join(c for c in testo if not unicodedata.combining(c))
    testo = re.sub(r"[^a-z0-9]+", "_", testo)
    return testo.strip("_")


def trova_colonna(df, possibili_nomi):
    colonne = {normalizza_testo(colonna): colonna for colonna in df.columns}

    for nome in possibili_nomi:
        normalizzato = normalizza_testo(nome)
        if normalizzato in colonne:
            return colonne[normalizzato]

    return None


def riconosci_colonne(df):
    return {
        "quantita": trova_colonna(
            df,
            ["quantita", "quantità", "quantity", "qty", "pezzi", "units", "unita", "unità"],
        ),
        "prezzo": trova_colonna(
            df,
            [
                "prezzo",
                "price",
                "prezzo_unitario",
                "unit_price",
                "unitprice",
                "prezzo unitario",
                "unit price",
            ],
        ),
        "prodotto": trova_colonna(
            df,
            ["prodotto", "product", "articolo", "item", "descrizione", "product_name"],
        ),
        "categoria": trova_colonna(
            df,
            ["categoria", "category", "reparto", "department", "segmento"],
        ),
        "data": trova_colonna(
            df,
            [
                "data",
                "date",
                "data_vendita",
                "sale_date",
                "data ordine",
                "order_date",
                "transaction_date",
            ],
        ),
        "cliente": trova_colonna(
            df,
            ["cliente", "customer", "client", "nome_cliente", "customer_name"],
        ),
    }


# ============================================================
# PULIZIA DATI
# ============================================================

def pulisci_dati(df):
    df = df.copy()
    righe_originali = len(df)
    colonne_originali = len(df.columns)

    df.dropna(how="all", inplace=True)
    df.dropna(axis=1, how="all", inplace=True)

    nomi = []
    contatori = {}
    for colonna in df.columns:
        nome = str(colonna).strip() or "Colonna"
        if nome in contatori:
            contatori[nome] += 1
            nome = f"{nome}_{contatori[nome]}"
        else:
            contatori[nome] = 1
        nomi.append(nome)
    df.columns = nomi

    colonne_testo = df.select_dtypes(include=["object", "string"]).columns
    for colonna in colonne_testo:
        df[colonna] = df[colonna].apply(
            lambda valore: valore.strip() if isinstance(valore, str) else valore
        )
        df[colonna] = df[colonna].replace(r"^\s*$", pd.NA, regex=True)

    duplicati = int(df.duplicated().sum())
    df.drop_duplicates(inplace=True)
    df.reset_index(drop=True, inplace=True)

    informazioni = {
        "Righe originali": righe_originali,
        "Righe finali": len(df),
        "Righe vuote rimosse": righe_originali - len(df) - duplicati,
        "Duplicati rimossi": duplicati,
        "Celle vuote": int(df.isna().sum().sum()),
        "Colonne originali": colonne_originali,
        "Colonne finali": len(df.columns),
    }
    return df, informazioni


# ============================================================
# CONVERSIONI ROBUSTE
# ============================================================

def converti_numero(valore):
    if pd.isna(valore):
        return None

    if isinstance(valore, bool):
        return float(valore)

    if isinstance(valore, (int, float)):
        return float(valore)

    testo = str(valore).strip()
    if not testo:
        return None

    negativo = testo.startswith("(") and testo.endswith(")")
    if negativo:
        testo = testo[1:-1]

    testo = (
        testo.replace("\u00a0", "")
        .replace("\u202f", "")
        .replace("€", "")
        .replace("$", "")
        .replace("£", "")
        .replace("CHF", "")
        .replace("EUR", "")
        .replace("USD", "")
        .replace("GBP", "")
        .replace("%", "")
        .replace("'", "")
        .replace(" ", "")
    )

    testo = re.sub(r"[^0-9,\.\-\+]", "", testo)
    if not testo or testo in {"-", "+", ".", ","}:
        return None

    if "," in testo and "." in testo:
        if testo.rfind(",") > testo.rfind("."):
            testo = testo.replace(".", "").replace(",", ".")
        else:
            testo = testo.replace(",", "")
    elif "," in testo:
        parti = testo.split(",")
        if len(parti) > 2:
            if all(len(p) == 3 for p in parti[1:]):
                testo = "".join(parti)
            else:
                testo = "".join(parti[:-1]) + "." + parti[-1]
        else:
            sinistra, destra = parti
            if len(destra) == 3 and sinistra not in {"0", "-0", "+0"}:
                testo = sinistra + destra
            else:
                testo = sinistra + "." + destra
    elif "." in testo:
        parti = testo.split(".")
        if len(parti) > 2 and all(len(p) == 3 for p in parti[1:]):
            testo = "".join(parti)
        elif len(parti) == 2 and len(parti[1]) == 3 and parti[0] not in {"0", "-0", "+0"}:
            testo = "".join(parti)

    try:
        numero = float(testo)
        return -numero if negativo and numero > 0 else numero
    except (ValueError, TypeError):
        return None


def converti_data(serie):
    originale = serie.copy()

    # Primo tentativo: formato europeo, molto comune nei file italiani.
    convertita = pd.to_datetime(originale, errors="coerce", dayfirst=True)

    # Secondo tentativo per i valori rimasti non riconosciuti.
    mancanti = convertita.isna() & originale.notna()
    if mancanti.any():
        secondo = pd.to_datetime(originale[mancanti], errors="coerce", dayfirst=False)
        convertita.loc[mancanti] = secondo

    return convertita


# ============================================================
# ANALISI VENDITE
# ============================================================

def analizza_vendite(df):
    df = df.copy()
    colonne = riconosci_colonne(df)

    quantita = colonne["quantita"]
    prezzo = colonne["prezzo"]
    prodotto = colonne["prodotto"]
    categoria = colonne["categoria"]
    data = colonne["data"]
    cliente = colonne["cliente"]

    analisi = {
        "attiva": False,
        "colonne": colonne,
        "kpi": {},
        "prodotti": pd.DataFrame(),
        "categorie": pd.DataFrame(),
        "clienti": pd.DataFrame(),
        "temporale": pd.DataFrame(),
        "anomalie": pd.DataFrame(),
        "metriche_qualita": {},
    }

    if not quantita or not prezzo:
        return df, analisi

    quantita_originale = df[quantita].copy()
    prezzo_originale = df[prezzo].copy()

    df[quantita] = df[quantita].apply(converti_numero)
    df[prezzo] = df[prezzo].apply(converti_numero)

    q_non_validi = int((quantita_originale.notna() & df[quantita].isna()).sum())
    p_non_validi = int((prezzo_originale.notna() & df[prezzo].isna()).sum())

    if data:
        df[data] = converti_data(df[data])
        date_non_valide = int(df[data].isna().sum())
    else:
        date_non_valide = 0

    df["Fatturato"] = df[quantita] * df[prezzo]

    maschera_valida = (
        df[quantita].notna()
        & df[prezzo].notna()
        & df["Fatturato"].notna()
    )
    valide = df.loc[maschera_valida].copy()

    if valide.empty:
        return df, analisi

    anomalie_mask = (
        df[quantita].isna()
        | df[prezzo].isna()
        | (df[quantita] < 0)
        | (df[prezzo] < 0)
    )
    if data:
        anomalie_mask = anomalie_mask | df[data].isna()

    anomalie = df.loc[anomalie_mask].copy()
    if not anomalie.empty:
        motivi = []
        for _, riga in anomalie.iterrows():
            problemi = []
            if pd.isna(riga[quantita]):
                problemi.append("Quantità non valida")
            elif riga[quantita] < 0:
                problemi.append("Quantità negativa")
            if pd.isna(riga[prezzo]):
                problemi.append("Prezzo non valido")
            elif riga[prezzo] < 0:
                problemi.append("Prezzo negativo")
            if data and pd.isna(riga[data]):
                problemi.append("Data non valida/mancante")
            motivi.append("; ".join(problemi))
        anomalie.insert(0, "Problema rilevato", motivi)

    analisi["anomalie"] = anomalie
    analisi["attiva"] = True

    fatturato_totale = float(valide["Fatturato"].sum())
    quantita_totale = float(valide[quantita].sum())
    numero_operazioni = int(len(valide))
    valore_medio = fatturato_totale / numero_operazioni if numero_operazioni else 0

    kpi = {
        "Fatturato totale": fatturato_totale,
        "Quantità totale": quantita_totale,
        "Numero operazioni": numero_operazioni,
        "Valore medio operazione": valore_medio,
    }

    if prodotto:
        prodotti = (
            valide.groupby(prodotto, dropna=False)
            .agg(
                Quantita_totale=(quantita, "sum"),
                Fatturato=("Fatturato", "sum"),
                Operazioni=("Fatturato", "size"),
            )
            .reset_index()
            .sort_values("Fatturato", ascending=False)
        )
        analisi["prodotti"] = prodotti
        if not prodotti.empty:
            kpi["Prodotto con maggior fatturato"] = str(prodotti.iloc[0][prodotto])

    if categoria:
        categorie = (
            valide.groupby(categoria, dropna=False)
            .agg(
                Quantita_totale=(quantita, "sum"),
                Fatturato=("Fatturato", "sum"),
                Operazioni=("Fatturato", "size"),
            )
            .reset_index()
            .sort_values("Fatturato", ascending=False)
        )
        analisi["categorie"] = categorie
        if not categorie.empty:
            kpi["Categoria migliore"] = str(categorie.iloc[0][categoria])

    if cliente:
        clienti = (
            valide.groupby(cliente, dropna=False)
            .agg(
                Fatturato=("Fatturato", "sum"),
                Operazioni=("Fatturato", "size"),
            )
            .reset_index()
            .sort_values("Fatturato", ascending=False)
        )
        analisi["clienti"] = clienti
        kpi["Clienti unici"] = int(valide[cliente].nunique(dropna=True))

    if data:
        temporale_base = valide.dropna(subset=[data]).copy()
        if not temporale_base.empty:
            temporale_base["Mese"] = temporale_base[data].dt.to_period("M").dt.to_timestamp()
            temporale = (
                temporale_base.groupby("Mese")
                .agg(
                    Fatturato=("Fatturato", "sum"),
                    Operazioni=("Fatturato", "size"),
                    Quantita_totale=(quantita, "sum"),
                )
                .reset_index()
                .sort_values("Mese")
            )
            analisi["temporale"] = temporale

            kpi["Prima vendita"] = temporale_base[data].min()
            kpi["Ultima vendita"] = temporale_base[data].max()

            if len(temporale) >= 2:
                precedente = float(temporale.iloc[-2]["Fatturato"])
                corrente = float(temporale.iloc[-1]["Fatturato"])
                if precedente != 0:
                    kpi["Variazione ultimo mese %"] = ((corrente - precedente) / abs(precedente)) * 100

    analisi["metriche_qualita"] = {
        "Valori quantità non convertibili": q_non_validi,
        "Valori prezzo non convertibili": p_non_validi,
        "Date non valide/mancanti": date_non_valide,
        "Righe valide per KPI": numero_operazioni,
        "Righe con anomalie": int(len(anomalie)),
    }

    analisi["kpi"] = kpi
    return df, analisi


# ============================================================
# STATISTICHE E QUALITÀ
# ============================================================

def crea_statistiche(df):
    risultati = []
    colonne_numeriche = df.select_dtypes(include="number").columns

    for colonna in colonne_numeriche:
        serie = df[colonna].dropna()
        if serie.empty:
            continue

        risultati.append(
            {
                "Colonna": colonna,
                "Valori": int(serie.count()),
                "Valori mancanti": int(df[colonna].isna().sum()),
                "Totale": serie.sum(),
                "Media": serie.mean(),
                "Mediana": serie.median(),
                "Deviazione standard": serie.std(),
                "Minimo": serie.min(),
                "Massimo": serie.max(),
            }
        )

    return pd.DataFrame(risultati)


def crea_analisi_qualita(df):
    risultati = []

    for colonna in df.columns:
        valori_mancanti = int(df[colonna].isna().sum())
        percentuale = (valori_mancanti / len(df) * 100) if len(df) else 0

        risultati.append(
            {
                "Colonna": colonna,
                "Tipo": str(df[colonna].dtype),
                "Valori mancanti": valori_mancanti,
                "% mancanti": round(percentuale, 2),
                "Valori unici": int(df[colonna].nunique(dropna=True)),
            }
        )

    return pd.DataFrame(risultati)


# ============================================================
# FORMATTAZIONE EXCEL
# ============================================================

def formatta_foglio(ws):
    riempimento_header = PatternFill("solid", fgColor=COLORE_PRINCIPALE)
    font_header = Font(color=COLORE_BIANCO, bold=True)
    bordo = Side(style="thin", color="D9E2F3")

    for cella in ws[1]:
        cella.fill = riempimento_header
        cella.font = font_header
        cella.alignment = Alignment(horizontal="center", vertical="center")

    for riga in ws.iter_rows():
        for cella in riga:
            cella.border = Border(bottom=bordo)
            cella.alignment = Alignment(vertical="center")

    for colonna in ws.columns:
        lettera = get_column_letter(colonna[0].column)
        lunghezza_massima = 0
        for cella in colonna[:500]:
            if cella.value is not None:
                lunghezza_massima = max(lunghezza_massima, len(str(cella.value)))
        ws.column_dimensions[lettera].width = min(lunghezza_massima + 3, 38)

    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def applica_formati_dati(ws):
    intestazioni = {cella.value: cella.column for cella in ws[1]}

    for nome, colonna in intestazioni.items():
        nome_norm = normalizza_testo(nome)
        for riga in range(2, ws.max_row + 1):
            cella = ws.cell(riga, colonna)
            if nome == "Fatturato" or "prezzo" in nome_norm or "price" in nome_norm:
                if isinstance(cella.value, (int, float)):
                    cella.number_format = '€ #,##0.00'
            elif "data" in nome_norm or "date" in nome_norm or nome == "Mese":
                if isinstance(cella.value, datetime):
                    cella.number_format = "dd/mm/yyyy"


# ============================================================
# DASHBOARD
# ============================================================

def crea_card(ws, titolo_range, valore_range, titolo, valore, formato=None, colore=None):
    ws.merge_cells(titolo_range)
    ws.merge_cells(valore_range)

    titolo_cell = ws[titolo_range.split(":")[0]]
    valore_cell = ws[valore_range.split(":")[0]]

    titolo_cell.value = titolo
    titolo_cell.font = Font(bold=True, color=COLORE_BIANCO, size=10)
    titolo_cell.fill = PatternFill("solid", fgColor=colore or COLORE_SECONDARIO)
    titolo_cell.alignment = Alignment(horizontal="center", vertical="center")

    valore_cell.value = valore
    valore_cell.font = Font(bold=True, size=17, color=COLORE_PRINCIPALE)
    valore_cell.fill = PatternFill("solid", fgColor=COLORE_CHIARO)
    valore_cell.alignment = Alignment(horizontal="center", vertical="center")

    if formato:
        valore_cell.number_format = formato


def crea_dashboard(workbook, analisi):
    if not analisi["attiva"]:
        return

    ws = workbook.create_sheet("Dashboard", 0)
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:L2")
    ws["A1"] = f"SMART SALES REPORT — {VERSION}"
    ws["A1"].font = Font(bold=True, size=24, color=COLORE_BIANCO)
    ws["A1"].fill = PatternFill("solid", fgColor=COLORE_PRINCIPALE)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    kpi = analisi["kpi"]

    crea_card(
        ws, "A4:C4", "A5:C6", "FATTURATO TOTALE",
        kpi.get("Fatturato totale", 0), '€ #,##0.00'
    )
    crea_card(
        ws, "D4:F4", "D5:F6", "QUANTITÀ VENDUTA",
        kpi.get("Quantità totale", 0), '#,##0.00'
    )
    crea_card(
        ws, "G4:I4", "G5:I6", "OPERAZIONI VALIDE",
        kpi.get("Numero operazioni", 0), '#,##0'
    )
    crea_card(
        ws, "J4:L4", "J5:L6", "VALORE MEDIO",
        kpi.get("Valore medio operazione", 0), '€ #,##0.00'
    )

    ws.merge_cells("A8:F8")
    ws["A8"] = f"Top prodotto: {kpi.get('Prodotto con maggior fatturato', 'N/D')}"
    ws["A8"].fill = PatternFill("solid", fgColor=COLORE_VERDE)
    ws["A8"].font = Font(bold=True)
    ws["A8"].alignment = Alignment(horizontal="center")

    ws.merge_cells("G8:L8")
    ws["G8"] = f"Top categoria: {kpi.get('Categoria migliore', 'N/D')}"
    ws["G8"].fill = PatternFill("solid", fgColor=COLORE_GIALLO)
    ws["G8"].font = Font(bold=True)
    ws["G8"].alignment = Alignment(horizontal="center")

    prodotti = analisi["prodotti"].head(10)
    if not prodotti.empty:
        start = 45
        nome_col = prodotti.columns[0]
        ws.cell(start, 1, nome_col)
        ws.cell(start, 2, "Fatturato")
        for n, (_, riga) in enumerate(prodotti.iterrows(), start=start + 1):
            ws.cell(n, 1, str(riga[nome_col]))
            ws.cell(n, 2, float(riga["Fatturato"]))

        grafico = BarChart()
        grafico.type = "col"
        grafico.style = 10
        grafico.title = "Top 10 prodotti per fatturato"
        grafico.y_axis.title = "Fatturato (€)"
        grafico.height = 10
        grafico.width = 16
        dati = Reference(ws, min_col=2, min_row=start, max_row=start + len(prodotti))
        categorie = Reference(ws, min_col=1, min_row=start + 1, max_row=start + len(prodotti))
        grafico.add_data(dati, titles_from_data=True)
        grafico.set_categories(categorie)
        grafico.legend = None
        ws.add_chart(grafico, "A10")

    categorie_df = analisi["categorie"].head(8)
    if not categorie_df.empty:
        start = 45
        nome_col = categorie_df.columns[0]
        ws.cell(start, 4, nome_col)
        ws.cell(start, 5, "Fatturato")
        for n, (_, riga) in enumerate(categorie_df.iterrows(), start=start + 1):
            ws.cell(n, 4, str(riga[nome_col]))
            ws.cell(n, 5, float(riga["Fatturato"]))

        grafico = PieChart()
        grafico.title = "Distribuzione fatturato per categoria"
        grafico.height = 10
        grafico.width = 14
        dati = Reference(ws, min_col=5, min_row=start, max_row=start + len(categorie_df))
        categorie = Reference(ws, min_col=4, min_row=start + 1, max_row=start + len(categorie_df))
        grafico.add_data(dati, titles_from_data=True)
        grafico.set_categories(categorie)
        grafico.dataLabels = DataLabelList()
        grafico.dataLabels.showPercent = True
        grafico.dataLabels.showLeaderLines = True
        ws.add_chart(grafico, "G10")

    temporale = analisi["temporale"]
    if not temporale.empty:
        start = 45
        ws.cell(start, 7, "Mese")
        ws.cell(start, 8, "Fatturato")
        for n, (_, riga) in enumerate(temporale.iterrows(), start=start + 1):
            ws.cell(n, 7, riga["Mese"].to_pydatetime() if hasattr(riga["Mese"], "to_pydatetime") else riga["Mese"])
            ws.cell(n, 8, float(riga["Fatturato"]))
            ws.cell(n, 7).number_format = "mmm yyyy"

        grafico = LineChart()
        grafico.title = "Andamento mensile del fatturato"
        grafico.y_axis.title = "Fatturato (€)"
        grafico.x_axis.title = "Mese"
        grafico.height = 9
        grafico.width = 22
        dati = Reference(ws, min_col=8, min_row=start, max_row=start + len(temporale))
        date = Reference(ws, min_col=7, min_row=start + 1, max_row=start + len(temporale))
        grafico.add_data(dati, titles_from_data=True)
        grafico.set_categories(date)
        grafico.legend = None
        ws.add_chart(grafico, "A29")

    for riga in range(45, 200):
        ws.row_dimensions[riga].hidden = True

    for colonna in range(1, 13):
        ws.column_dimensions[get_column_letter(colonna)].width = 13


# ============================================================
# GRAFICI NEI FOGLI DI ANALISI
# ============================================================

def aggiungi_grafico_prodotti(workbook):
    if "Vendite per Prodotto" not in workbook.sheetnames:
        return

    ws = workbook["Vendite per Prodotto"]
    if ws.max_row < 2:
        return

    grafico = BarChart()
    grafico.type = "col"
    grafico.style = 10
    grafico.title = "Fatturato per prodotto"
    grafico.y_axis.title = "Fatturato (€)"
    grafico.x_axis.title = "Prodotto"
    grafico.height = 11
    grafico.width = 18

    dati = Reference(ws, min_col=3, min_row=1, max_row=ws.max_row)
    categorie = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    grafico.add_data(dati, titles_from_data=True)
    grafico.set_categories(categorie)
    grafico.legend = None
    ws.add_chart(grafico, "F4")


def aggiungi_grafico_categorie(workbook):
    if "Vendite per Categoria" not in workbook.sheetnames:
        return

    ws = workbook["Vendite per Categoria"]
    if ws.max_row < 2:
        return

    grafico = PieChart()
    grafico.title = "Distribuzione fatturato per categoria"
    grafico.height = 11
    grafico.width = 16

    dati = Reference(ws, min_col=3, min_row=1, max_row=ws.max_row)
    categorie = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    grafico.add_data(dati, titles_from_data=True)
    grafico.set_categories(categorie)
    grafico.dataLabels = DataLabelList()
    grafico.dataLabels.showPercent = True
    grafico.dataLabels.showLeaderLines = True
    ws.add_chart(grafico, "F4")


def aggiungi_grafico_temporale(workbook):
    if "Andamento Mensile" not in workbook.sheetnames:
        return

    ws = workbook["Andamento Mensile"]
    if ws.max_row < 2:
        return

    grafico = LineChart()
    grafico.title = "Andamento mensile del fatturato"
    grafico.y_axis.title = "Fatturato (€)"
    grafico.x_axis.title = "Mese"
    grafico.height = 11
    grafico.width = 20

    dati = Reference(ws, min_col=2, min_row=1, max_row=ws.max_row)
    categorie = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    grafico.add_data(dati, titles_from_data=True)
    grafico.set_categories(categorie)
    grafico.legend = None
    ws.add_chart(grafico, "F4")


# ============================================================
# CREAZIONE REPORT
# ============================================================

def genera_report(df, statistiche, qualita, pulizia, analisi, percorso_originale):
    cartella_output = percorso_originale.parent / "output"
    cartella_output.mkdir(exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    nome_file = f"REPORT_{percorso_originale.stem}_{timestamp}.xlsx"
    percorso_output = cartella_output / nome_file

    info_righe = [
        ("Applicazione", APP_NAME),
        ("Versione", VERSION),
        ("File originale", percorso_originale.name),
        ("Data generazione", datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
    ]
    info_righe.extend(list(pulizia.items()))
    info_righe.append(("Analisi vendite", "Disponibile" if analisi["attiva"] else "Non rilevata"))
    info_righe.extend(list(analisi.get("metriche_qualita", {}).items()))
    informazioni = pd.DataFrame(info_righe, columns=["Informazione", "Valore"])

    with pd.ExcelWriter(percorso_output, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Dati Puliti", index=False)
        statistiche.to_excel(writer, sheet_name="Statistiche", index=False)
        qualita.to_excel(writer, sheet_name="Qualita Dati", index=False)
        informazioni.to_excel(writer, sheet_name="Informazioni", index=False)

        if analisi["attiva"] and not analisi["prodotti"].empty:
            analisi["prodotti"].to_excel(writer, sheet_name="Vendite per Prodotto", index=False)

        if analisi["attiva"] and not analisi["categorie"].empty:
            analisi["categorie"].to_excel(writer, sheet_name="Vendite per Categoria", index=False)

        if analisi["attiva"] and not analisi["clienti"].empty:
            analisi["clienti"].to_excel(writer, sheet_name="Vendite per Cliente", index=False)

        if analisi["attiva"] and not analisi["temporale"].empty:
            analisi["temporale"].to_excel(writer, sheet_name="Andamento Mensile", index=False)

        if analisi["attiva"] and not analisi["anomalie"].empty:
            analisi["anomalie"].to_excel(writer, sheet_name="Anomalie", index=False)

        workbook = writer.book
        crea_dashboard(workbook, analisi)

        for nome_foglio in workbook.sheetnames:
            if nome_foglio != "Dashboard":
                formatta_foglio(workbook[nome_foglio])
                applica_formati_dati(workbook[nome_foglio])

        if "Qualita Dati" in workbook.sheetnames:
            ws_qualita = workbook["Qualita Dati"]
            if ws_qualita.max_row >= 2:
                ws_qualita.conditional_formatting.add(
                    f"D2:D{ws_qualita.max_row}",
                    ColorScaleRule(
                        start_type="min",
                        start_color="63BE7B",
                        mid_type="percentile",
                        mid_value=50,
                        mid_color="FFEB84",
                        end_type="max",
                        end_color="F8696B",
                    ),
                )

        if "Anomalie" in workbook.sheetnames:
            ws_anomalie = workbook["Anomalie"]
            for cella in ws_anomalie[1]:
                cella.fill = PatternFill("solid", fgColor="C00000")
                cella.font = Font(color=COLORE_BIANCO, bold=True)

        aggiungi_grafico_prodotti(workbook)
        aggiungi_grafico_categorie(workbook)
        aggiungi_grafico_temporale(workbook)

    logger.info("Report creato: %s", percorso_output)
    return percorso_output


# ============================================================
# APERTURA REPORT
# ============================================================

def apri_report(percorso):
    try:
        if os.name == "nt":
            os.startfile(percorso)
    except Exception as errore:
        logger.warning("Impossibile aprire automaticamente il report: %s", errore)


# ============================================================
# PROGRAMMA PRINCIPALE
# ============================================================

def main():
    logger.info("Avvio %s versione %s", APP_NAME, VERSION)

    print("=" * 65)
    print(f"{APP_NAME} - Versione {VERSION}")
    print("=" * 65)

    percorso = scegli_file()
    if percorso is None:
        print("Nessun file selezionato.")
        return

    try:
        print(f"\nFile selezionato: {percorso.name}")
        logger.info("File selezionato: %s", percorso)

        print("\nCaricamento dati...")
        df_originale = carica_dati(percorso)

        if df_originale.empty:
            raise ValueError("Il file selezionato è vuoto.")

        print("Pulizia dati...")
        df, pulizia = pulisci_dati(df_originale)

        if df.empty:
            raise ValueError("Dopo la pulizia non sono rimaste righe utilizzabili.")

        print("Analisi dati...")
        df, analisi = analizza_vendite(df)

        statistiche = crea_statistiche(df)
        qualita = crea_analisi_qualita(df)

        print("\n=== PULIZIA DATI ===")
        for nome, valore in pulizia.items():
            print(f"{nome}: {valore}")

        if analisi["attiva"]:
            print("\n=== ANALISI VENDITE ===")
            for nome, valore in analisi["kpi"].items():
                if isinstance(valore, float):
                    print(f"{nome}: {valore:,.2f}")
                elif isinstance(valore, pd.Timestamp):
                    print(f"{nome}: {valore.strftime('%d/%m/%Y')}")
                else:
                    print(f"{nome}: {valore}")

            anomalie = len(analisi["anomalie"])
            if anomalie:
                print(f"\nAttenzione: rilevate {anomalie} righe da controllare.")
        else:
            print(
                "\nIl file è stato pulito, ma non sono state riconosciute "
                "colonne sufficienti per l'analisi vendite (quantità + prezzo)."
            )

        print("\nGenerazione report Excel...")
        output = genera_report(
            df,
            statistiche,
            qualita,
            pulizia,
            analisi,
            percorso,
        )

        print("\n" + "=" * 65)
        print("REPORT CREATO CON SUCCESSO")
        print("=" * 65)
        print(f"\nFile salvato in:\n{output}")

        root = crea_root()
        messagebox.showinfo(
            "Report completato",
            f"Analisi completata con successo.\n\nReport salvato in:\n{output}",
        )
        root.destroy()

        apri_report(output)

    except PermissionError:
        logger.exception("Permesso negato durante la scrittura del report.")
        messaggio = (
            "Impossibile scrivere il file.\n"
            "Chiudi eventuali report Excel già aperti e riprova."
        )
        print(f"\nERRORE: {messaggio}")

        root = crea_root()
        messagebox.showerror("Errore", messaggio)
        root.destroy()

    except Exception as errore:
        logger.exception("Errore durante l'elaborazione: %s", errore)
        print("\nERRORE DURANTE L'ELABORAZIONE:")
        print(str(errore))

        root = crea_root()
        messagebox.showerror("Errore", str(errore))
        root.destroy()


if __name__ == "__main__":
    main()
